"""Regression coverage for data, prediction, persistence, CLI, and Tk events."""

from contextlib import redirect_stderr, redirect_stdout
from concurrent.futures import Future
import io
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

import numpy as np
from docx import Document

from word_predict import (
    LSTMModel,
    NextWordPredictor,
    NextWordPredictorGUI,
    PAD_TOKEN,
    TextDataset,
    build_parser,
    create_and_train_model,
    main,
    tf,
)


class TemporaryCorpusTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.corpus = self.folder / "corpus.txt"
        self.corpus.write_text("One two three one two four", encoding="utf-8")


class DatasetTests(TemporaryCorpusTest):
    def test_sequences_and_targets_use_same_vocabulary(self):
        dataset = TextDataset(self.corpus, 2)
        inputs, targets = dataset.prepare_training_data()
        self.assertEqual(dataset.vocabulary, [PAD_TOKEN, "four", "one", "three", "two"])
        decoded = [[dataset.index_to_word[int(index)] for index in row] for row in inputs]
        self.assertEqual(decoded, [["one", "two"], ["two", "three"], ["three", "one"], ["one", "two"]])
        self.assertEqual([dataset.index_to_word[int(index)] for index in targets], ["three", "one", "two", "four"])
        self.assertEqual(inputs.dtype, np.int32)

    def test_bom_and_windows_encoding(self):
        self.corpus.write_text("One two three", encoding="utf-8-sig")
        self.assertEqual(TextDataset(self.corpus, 1).words, ["one", "two", "three"])
        self.corpus.write_bytes("One \u201ctwo\u201d three".encode("cp1252"))
        self.assertEqual(TextDataset(self.corpus, 1).words, ["one", "\u201ctwo\u201d", "three"])

    def test_docx_accepts_path_and_uppercase_extension(self):
        path = self.folder / "corpus.DOCX"
        document = Document()
        document.add_paragraph("One TWO")
        document.add_paragraph("Three four")
        document.save(path)
        self.assertEqual(TextDataset(path, 2).words, ["one", "two", "three", "four"])

    def test_rejects_invalid_training_data(self):
        for text in ("", "one", "one two"):
            with self.subTest(text=text):
                self.corpus.write_text(text, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "at least 3 words"):
                    TextDataset(self.corpus, 2)
        with self.assertRaisesRegex(ValueError, "positive"):
            TextDataset(self.corpus, 0)
        with self.assertRaisesRegex(ValueError, "txt or .docx"):
            TextDataset(self.folder / "data.csv", 2)
        with self.assertRaises(FileNotFoundError):
            TextDataset(self.folder / "missing.txt", 2)


class ScoreModel:
    vocabulary = [PAD_TOKEN, "one", "two", "three", "four"]
    sequence_length = 3

    def __init__(self):
        self.inputs = []

    def __call__(self, values, training=False):
        self.inputs.append(values)
        return np.array([[0.5, 0.05, 0.1, 0.2, 0.15]])


class PredictionTests(unittest.TestCase):
    def setUp(self):
        self.model = ScoreModel()
        self.predictor = NextWordPredictor(self.model)

    def test_top_three_exclude_padding_and_preserve_probabilities(self):
        self.assertEqual(self.predictor.predict_next_word("ONE"),
                         [("three", 0.2), ("four", 0.15), ("two", 0.1)])
        np.testing.assert_array_equal(self.model.inputs[-1], [[1, 0, 0]])

    def test_empty_and_unknown_input_do_not_run_inference(self):
        for text in ("", "  \n", "unseen words"):
            self.assertEqual(self.predictor.predict_next_word(text), [])
        self.assertEqual(self.model.inputs, [])

    def test_recent_context_and_unknown_words(self):
        self.predictor.predict_next_word("one two unknown three")
        np.testing.assert_array_equal(self.model.inputs[-1], [[2, 3, 0]])

    def test_top_k_larger_than_vocabulary_is_safe(self):
        self.assertEqual(len(self.predictor.predict_next_word("one", top_k=20)), 4)
        with self.assertRaises(ValueError):
            self.predictor.predict_next_word("one", top_k=0)

    def test_legacy_model_requests_retraining(self):
        with self.assertRaisesRegex(ValueError, "--retrain"):
            NextWordPredictor(LSTMModel(vocab_size=5))

    def test_invalid_scores_are_reported(self):
        for scores in (np.zeros((2, 5)), np.full((1, 5), np.nan)):
            with self.subTest(shape=scores.shape):
                with patch.object(ScoreModel, "__call__", return_value=scores):
                    with self.assertRaisesRegex(ValueError, "invalid prediction"):
                        self.predictor.predict_next_word("one")


class ModelAndCLITests(TemporaryCorpusTest):
    def test_train_save_load_preserves_predictions_without_corpus(self):
        tf.keras.utils.set_random_seed(42)
        self.corpus.write_text("one two three one two four " * 3, encoding="utf-8")
        dataset = TextDataset(self.corpus, 2)
        model = create_and_train_model(dataset, epochs=1, batch_size=2,
                                       embedding_dim=4, hidden_size=8, num_layers=2)
        context = np.array([[dataset.word_to_index["one"], 0]], dtype=np.int32)
        expected = model(context, training=False).numpy()
        model_path = self.folder / "model.keras"
        model.save(model_path)
        self.corpus.unlink()
        restored = tf.keras.models.load_model(model_path, compile=False)
        self.assertEqual(list(restored.vocabulary), dataset.vocabulary)
        self.assertEqual(restored.sequence_length, 2)
        np.testing.assert_allclose(restored(context, training=False), expected, atol=1e-6)
        self.assertEqual(len(NextWordPredictor(restored).predict_next_word("one")), 3)
        output = io.StringIO()
        with redirect_stdout(output):
            status = main(["--model-path", str(model_path), "--text-file", str(self.corpus),
                           "--sequence-length", "99", "--predict", "one"])
        self.assertEqual(status, 0)
        self.assertEqual(output.getvalue().count("%"), 3)

    def test_cli_can_train_one_example_create_directories_and_retrain(self):
        path = self.folder / "nested" / "model.keras"
        options = ["--text-file", str(self.corpus), "--model-path", str(path),
                   "--sequence-length", "5", "--epochs", "1", "--batch-size", "2",
                   "--embedding-dim", "4", "--hidden-size", "4", "--num-layers", "1",
                   "--train-only"]
        self.assertEqual(main(options), 0)
        self.assertTrue(path.is_file())
        self.corpus.write_text("red green blue red green yellow", encoding="utf-8")
        self.assertEqual(main([*options, "--retrain"]), 0)
        restored = tf.keras.models.load_model(path, compile=False)
        self.assertIn("yellow", restored.vocabulary)
        self.assertNotIn("one", restored.vocabulary)

    def test_cli_errors_have_failure_exit_codes(self):
        with redirect_stderr(io.StringIO()):
            self.assertEqual(main(["--text-file", str(self.folder / "missing.txt"),
                                   "--model-path", str(self.folder / "model.keras"), "--train-only"]), 1)
            for options in (["--epochs", "0"], ["--sequence-length", "-1"],
                            ["--train-only", "--predict", "one"]):
                with self.assertRaises(SystemExit) as error:
                    build_parser().parse_args(options)
                self.assertEqual(error.exception.code, 2)
            with self.assertRaises(SystemExit) as error:
                main(["--model-path", "model.h5"])
            self.assertEqual(error.exception.code, 2)


class GUITests(unittest.TestCase):
    def setUp(self):
        import tkinter as tk

        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"Tk display unavailable: {error}")
        self.root.withdraw()
        self.app = NextWordPredictorGUI(self.root, NextWordPredictor(ScoreModel()))
        self.root.update()
        self.addCleanup(self.app.close)

    def complete_prediction(self, text, predictions):
        self.app.future = Future()
        self.app.future.set_result(predictions)
        self.app.poll_prediction(text)

    def test_clear_input_removes_buttons_without_crashing(self):
        self.app.input_text.set("one")
        self.complete_prediction("one", [("two", 0.5)])
        self.assertEqual(len(self.app.prediction_buttons), 1)
        self.app.input_text.set("")
        self.assertEqual(self.app.prediction_buttons, [])
        self.assertEqual(self.app.status.get(), "Waiting for input...")
        self.assertIsNone(self.app.debounce_id)

    def test_clicked_word_avoids_duplicate_spaces(self):
        self.app.input_text.set("one  ")
        self.complete_prediction("one  ", [("two", 0.5)])
        self.app.prediction_buttons[0].invoke()
        self.assertEqual(self.app.input_text.get(), "one two")
        self.assertIsNotNone(self.app.debounce_id)

    def test_stale_results_do_not_replace_current_text(self):
        self.app.input_text.set("three")
        self.complete_prediction("one", [("two", 0.5)])
        self.assertEqual(self.app.prediction_buttons, [])
        self.assertEqual(self.app.input_text.get(), "three")

    def test_pasted_input_runs_background_prediction_through_tk_events(self):
        self.app.input_text.set("one")
        deadline = time.monotonic() + 5
        while not self.app.prediction_buttons and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)
        self.assertEqual(len(self.app.prediction_buttons), 3)
        self.assertEqual(self.app.prediction_buttons[0].cget("text"), "three (20.0%)")


if __name__ == "__main__":
    unittest.main()
